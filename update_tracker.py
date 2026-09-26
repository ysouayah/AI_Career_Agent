"""
Copies new rows from agent_log.csv into agent_tracker.xlsx without touching your labels.

    python update_tracker.py                      # defaults below
    python update_tracker.py path/to/log.csv path/to/tracker.xlsx

A row is new if its (run date, URL) pair isn't in the tracker yet, so running it twice is safe.
Approved and near-miss jobs are marked Review? = "Yes". A stable ~10% of filtered and
low-scoring jobs are marked "Sample" -- the same URL is always in or out of the sample, so the
sample can't be nudged by re-running.
"""
import csv
import hashlib
import sys
from datetime import datetime
from openpyxl import load_workbook

LOG = sys.argv[1] if len(sys.argv) > 1 else "agent_log.csv"
TRACKER = sys.argv[2] if len(sys.argv) > 2 else "agent_tracker.xlsx"
SAMPLE_RATE = 0.10

# Tracker columns A..K, in order.
FIELDS = ["run_date", "company", "title", "location", "decision", "score", "reason",
          "hiring_track", "rules_version", "url"]


def review_flag(row):
    if row["decision"] in ("approved", "near_miss"):
        return "Yes"
    if row["decision"] in ("filtered", "below_floor") and row["url"]:
        bucket = int(hashlib.sha1(row["url"].encode()).hexdigest()[:8], 16) / 0xFFFFFFFF
        return "Sample" if bucket < SAMPLE_RATE else ""
    return ""


def main():
    wb = load_workbook(TRACKER)
    ws = wb["Agent Log"]

    existing, last = set(), 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=2).value:            # company filled = used row
            last = r
            run = ws.cell(row=r, column=1).value
            run = run.strftime("%Y-%m-%d") if hasattr(run, "strftime") else str(run or "")
            existing.add((run, str(ws.cell(row=r, column=10).value or "")))

    with open(LOG, newline="", encoding="utf-8") as f:
        rows = [row for row in csv.DictReader(f)
                if (row["run_date"], row["url"]) not in existing]

    for i, row in enumerate(rows, start=last + 1):
        for c, key in enumerate(FIELDS, start=1):
            value = row.get(key, "")
            if key == "run_date" and value:
                value = datetime.strptime(value, "%Y-%m-%d").date()
            elif key == "score" and value:
                value = int(float(value))
            ws.cell(row=i, column=c, value=value or None)
        ws.cell(row=i, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=i, column=11, value=review_flag(row) or None)

    wb.save(TRACKER)
    flagged = sum(1 for r in rows if review_flag(r))
    print(f"Added {len(rows)} row(s) to {TRACKER}; {flagged} flagged for review.")


if __name__ == "__main__":
    main()
